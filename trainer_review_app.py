"""
Streamlit app for trainers to review and edit evaluations.
Usage: streamlit run trainer_review_app.py --server.port 8002
"""

import streamlit as st
import os
import sys
import django
from pathlib import Path
import json
import zipfile
from io import BytesIO

# Fix Python path for nested directory structure
# Current: /DarkSpace/InternshipPortal/trainer_review_app.py
# Need:    /DarkSpace/InternshipPortal/InternshipPortal/InternshipPortal/settings.py
app_root = Path(__file__).parent  # /DarkSpace/InternshipPortal
project_root = app_root / "InternshipPortal"  # /DarkSpace/InternshipPortal/InternshipPortal

if str(project_root) not in sys.path:
    sys.path.insert(0, str(project_root))

# Use pymysql as MySQLdb (before Django setup)
import pymysql
pymysql.install_as_MySQLdb()
# Spoof version so Django's mysqlclient version check passes
pymysql.version_info = (2, 2, 1, "final", 0)

# Load the innermost .env (contains USE_S3, AWS settings) before Django reads it
from dotenv import load_dotenv
_env_path = project_root / "InternshipPortal" / ".env"
if _env_path.exists():
    load_dotenv(_env_path, override=True)

# Setup Django
os.environ.setdefault("DJANGO_SETTINGS_MODULE", "InternshipPortal.settings")
django.setup()

from django.conf import settings
from django.utils import timezone
from django.db.models import Min, Q
from django.contrib.auth.hashers import make_password, check_password
from Submissions.models import Evaluation, Student, Submission, TrainerProfile, TrainerEditLog

# ── File reading helpers ──────────────────────────────────────────────────────

_USE_S3 = os.getenv('USE_S3', 'False').lower() == 'true'

# All locations where media files may live (tried in order)
_MEDIA_ROOTS = [
    '/DarkSpace/InternshipPortal/media',
    '/home/ubuntu/InternshipPortal/media',
    '/DarkSpace/InternshipPortal/InternshipPortal/media',
    '/app/media',
]


def _resolve_local_path(stored_path: str):
    """
    Try to find a file that may be stored under any of the known media roots.
    stored_path can be:
      - absolute: /DarkSpace/InternshipPortal/media/Evaluations/...
      - relative: Evaluations/...
    Returns the first existing path, or None.
    """
    if os.path.isabs(stored_path):
        if os.path.exists(stored_path):
            return stored_path
        # Extract relative part after /media/
        if '/media/' in stored_path:
            rel = stored_path.split('/media/', 1)[1]
        else:
            return None
    else:
        rel = stored_path.lstrip('/')

    for root in _MEDIA_ROOTS:
        candidate = os.path.join(root, rel)
        if os.path.exists(candidate):
            return candidate
    return None


def _read_eval_file(stored_path: str) -> str:
    """Read an evaluation text file. Tries local paths then S3."""
    local = _resolve_local_path(stored_path)
    if local:
        try:
            with open(local, 'r', encoding='utf-8') as f:
                return f.read()
        except Exception as e:
            return f"Error reading file: {e}"

    if _USE_S3:
        rel = stored_path.split('/media/', 1)[1] if '/media/' in stored_path else stored_path.lstrip('/')
        try:
            from django.core.files.storage import default_storage
            with default_storage.open(rel, 'r') as f:
                return f.read()
        except Exception as e:
            return f"Error reading from S3: {e}"

    return f"Error: Evaluation file not found ({stored_path})"


def _open_submission_file(file_name: str):
    """
    Open a submission file as bytes. Tries S3 first (if enabled),
    then falls back to local media roots.
    """
    if _USE_S3:
        try:
            from django.core.files.storage import default_storage
            return default_storage.open(file_name, 'rb').read()
        except Exception:
            pass

    local = _resolve_local_path(file_name)
    if local:
        with open(local, 'rb') as f:
            return f.read()

    raise FileNotFoundError(f"Submission file not found: {file_name}")


# Import nbconvert for notebook rendering
try:
    from nbconvert import HTMLExporter
    NBCONVERT_AVAILABLE = True
except ImportError:
    NBCONVERT_AVAILABLE = False

# Try importing pptx and PIL for advanced file handling
try:
    from pptx import Presentation
    PPTX_AVAILABLE = True
except ImportError:
    PPTX_AVAILABLE = False

try:
    from PIL import Image
    PIL_AVAILABLE = True
except ImportError:
    PIL_AVAILABLE = False


# ============================================================================
# ZIP FILE PREVIEW HELPER
# ============================================================================

def render_zip_explorer(file_path):
    """Display an interactive zip file explorer with preview functionality."""
    try:
        with zipfile.ZipFile(file_path, 'r') as zip_ref:
            file_list = zip_ref.namelist()
            
            if not file_list:
                st.info("Zip file is empty")
                return
            
            st.write(f"**Zip Contents:** {len([f for f in file_list if not f.endswith('/')])} files")
            
            for file_name in sorted(file_list):
                # Skip directories
                if file_name.endswith('/'):
                    continue
                
                with st.expander(f"📄 {file_name}", expanded=False):
                    try:
                        file_data = zip_ref.read(file_name)
                        file_ext = Path(file_name).suffix.lower()
                        
                        # Preview based on file type
                        if file_ext in ['.png', '.jpg', '.jpeg', '.gif', '.bmp']:
                            if PIL_AVAILABLE:
                                image = Image.open(BytesIO(file_data))
                                st.image(image, caption=file_name, use_column_width=True)
                            else:
                                st.download_button(label="Download Image", data=file_data, file_name=file_name, key=f"zip_img_{file_name}")
                        
                        elif file_ext in ['.txt', '.sql', '.py', '.js', '.json', '.csv', '.xml', '.html']:
                            text_content = file_data.decode('utf-8', errors='ignore')
                            st.code(text_content, language='python' if file_ext == '.py' else 'sql' if file_ext == '.sql' else 'json' if file_ext == '.json' else 'text')
                            st.download_button(label="Download", data=file_data, file_name=file_name, key=f"zip_txt_{file_name}")
                        
                        elif file_ext == '.ipynb':
                            if NBCONVERT_AVAILABLE:
                                try:
                                    import tempfile
                                    with tempfile.NamedTemporaryFile(suffix='.ipynb', delete=False, mode='wb') as tmp:
                                        tmp.write(file_data)
                                        tmp.flush()
                                        
                                        exporter = HTMLExporter()
                                        try:
                                            # Try standard HTML export first
                                            html_content, _ = exporter.from_filename(tmp.name)
                                            st.components.v1.html(html_content, height=600, scrolling=True)
                                        except AttributeError as ae:
                                            # Workaround for nbconvert parser bugs
                                            if 'parse_axt_heading' in str(ae):
                                                import json
                                                with open(tmp.name, 'r', encoding='utf-8') as f:
                                                    nb = json.load(f)
                                                
                                                html_parts = ['<html><body><style>body{font-family:Arial;padding:20px;} .cell{margin:20px 0;border:1px solid #ddd;padding:10px;} .code{background:#f5f5f5;padding:10px;overflow-x:auto;} .markdown{background:#f9f9f9;padding:10px;}</style>']
                                                
                                                for cell in nb.get('cells', []):
                                                    if cell['cell_type'] == 'code':
                                                        source = ''.join(cell.get('source', []))
                                                        html_parts.append(f'<div class="cell"><div class="code"><pre><code>{source}</code></pre></div></div>')
                                                    elif cell['cell_type'] == 'markdown':
                                                        source = ''.join(cell.get('source', []))
                                                        html_parts.append(f'<div class="cell markdown"><div>{source}</div></div>')
                                                
                                                html_parts.append('</body></html>')
                                                st.components.v1.html(''.join(html_parts), height=600, scrolling=True)
                                            else:
                                                raise ae
                                    
                                    os.unlink(tmp.name)
                                except Exception as e:
                                    st.warning(f"Could not render notebook: {str(e)}")
                                    st.download_button(label="Download Notebook", data=file_data, file_name=file_name, key=f"zip_nb_{file_name}")
                            else:
                                st.download_button(label="Download Notebook", data=file_data, file_name=file_name, key=f"zip_nb_{file_name}")
                        
                        elif file_ext == '.pptx':
                            if PPTX_AVAILABLE:
                                try:
                                    prs = Presentation(BytesIO(file_data))
                                    st.write(f"**Slides:** {len(prs.slides)}")
                                    for i, slide in enumerate(prs.slides[:3]):
                                        st.write(f"Slide {i+1}:")
                                        for shape in slide.shapes:
                                            if hasattr(shape, "text") and shape.text:
                                                st.write(f"  {shape.text}")
                                except Exception as e:
                                    st.warning(f"Could not parse PPTX: {str(e)}")
                                st.download_button(label="Download PPTX", data=file_data, file_name=file_name, key=f"zip_pptx_{file_name}")
                            else:
                                st.download_button(label="Download", data=file_data, file_name=file_name, key=f"zip_ppt_{file_name}")
                        
                        elif file_ext == '.pdf':
                            st.download_button(label="Download PDF", data=file_data, file_name=file_name, key=f"zip_pdf_{file_name}")
                        
                        else:
                            st.write(f"File type: {file_ext}")
                            st.download_button(label="Download", data=file_data, file_name=file_name, key=f"zip_file_{file_name}")
                    
                    except Exception as e:
                        st.error(f"Error processing {file_name}: {str(e)}")
    
    except Exception as e:
        st.error(f"Error reading zip file: {str(e)}")

# Page config
st.set_page_config(page_title="Trainer Evaluation Review", page_icon="📋", layout="wide", initial_sidebar_state="expanded")
st.title("📋 Evaluation Review Portal")

# ============================================================================
# SESSION STATE MANAGEMENT
# ============================================================================

for _key in ["trainer_email", "trainer_name", "selected_eval_id"]:
    if _key not in st.session_state:
        st.session_state[_key] = None
if "trainer_is_admin" not in st.session_state:
    st.session_state["trainer_is_admin"] = False
for _key in ["eval_text_edited", "eval_grade_edited", "eval_text_original", "eval_edit_feedback"]:
    if _key not in st.session_state:
        st.session_state[_key] = {}

# ============================================================================
# TRAINER LOGIN
# ============================================================================

def trainer_login():
    st.header("Trainer Login")
    col1, col2, col3 = st.columns([1, 2, 1])
    with col2:
        trainer_email = st.text_input("Email", placeholder="trainer@example.com")
        trainer_password = st.text_input("Password", type="password")
        if st.button("Login", type="primary", use_container_width=True):
            trainer = TrainerProfile.objects.filter(email=trainer_email, is_active=True).first()
            if trainer and check_password(trainer_password, trainer.password_hash):
                st.session_state.trainer_email = trainer_email
                st.session_state.trainer_name = trainer.name
                st.session_state.trainer_is_admin = trainer.is_admin
                st.rerun()
            else:
                st.error("Invalid email or password.")

def trainer_logout():
    st.session_state.trainer_email = None
    st.session_state.trainer_name = None
    st.session_state.trainer_is_admin = False
    st.session_state.selected_eval_id = None
    st.rerun()

# ============================================================================
# MAIN APP
# ============================================================================

if st.session_state.trainer_email is None:
    trainer_login()
else:
    # ── Sidebar filters ───────────────────────────────────────────────────────
    with st.sidebar:
        display = st.session_state.trainer_name or st.session_state.trainer_email
        st.write(f"**{display}**")
        st.caption(st.session_state.trainer_email)
        if st.session_state.trainer_is_admin:
            st.caption("🔑 Admin")
        if st.button("Logout", use_container_width=True):
            trainer_logout()

        st.divider()
        st.subheader("Filters")

        selected_course = st.selectbox("Course", ["All", "CDA", "AIE", "CDE"])

        review_status = st.radio(
            "Review status",
            ["All (email not sent)", "Not yet reviewed", "Reviewed — email pending"],
            index=0,
        )

        submission_source = st.radio(
            "Submission source",
            ["All", "Portal only", "Google Sheets only"],
            index=0,
        )

        search_team    = st.text_input("Search Team ID",    placeholder="PTID-CDA-...")
        search_project = st.text_input("Search Project ID", placeholder="PRDA-02")

        st.divider()
        st.subheader("Submission date range")
        date_from = st.date_input("From", value=None)
        date_to   = st.date_input("To",   value=None)

        st.divider()
        pending_days_min = st.number_input(
            "Min days pending (eval age)", min_value=0, value=0, step=1,
            help="Only show evaluations created at least N days ago",
        )

        sort_order = st.selectbox(
            "Sort by",
            ["Oldest first", "Newest first", "Project ID", "Team ID"],
        )

    # ── Helper: earliest submission date for an evaluation ───────────────────
    def get_submission_date(eval_obj):
        sub = (
            Submission.objects
            .filter(
                team_project__team=eval_obj.team,
                team_project__project=eval_obj.project,
            )
            .order_by("created_at")
            .first()
        )
        return sub.created_at if sub else None

    # ── List view ─────────────────────────────────────────────────────────────
    if not st.session_state.selected_eval_id:
        _tab_labels = ["Pending Evaluations", "Instructions"]
        if st.session_state.trainer_is_admin:
            _tab_labels.append("Admin Dashboard")
        _tabs = st.tabs(_tab_labels)
        tab1 = _tabs[0]
        tab2 = _tabs[1]
        tab_admin = _tabs[2] if st.session_state.trainer_is_admin else None

        with tab1:
            st.subheader("Evaluations — Email Not Yet Sent")

            # Base: email not sent yet  OR  came from Google Sheets (audit trail)
            # Excludes CDS (auto-approved, no trainer action needed)
            base_qs = (
                Evaluation.objects
                .filter(Q(email_sent=False) | Q(sheets_submitted=True))
                .exclude(course="CDS")
                .select_related("team", "project")
            )

            if selected_course != "All":
                base_qs = base_qs.filter(course=selected_course)

            # Source filter narrows before review-status, so counts stay consistent
            if submission_source == "Portal only":
                base_qs = base_qs.filter(sheets_submitted=False, email_sent=False)
            elif submission_source == "Google Sheets only":
                base_qs = base_qs.filter(sheets_submitted=True)

            if review_status == "Not yet reviewed":
                base_qs = base_qs.filter(reviewed=False)
            elif review_status == "Reviewed — email pending":
                # Only meaningful for portal submissions where email hasn't fired yet
                base_qs = base_qs.filter(reviewed=True, email_sent=False)

            if search_team.strip():
                base_qs = base_qs.filter(team__team_id__icontains=search_team.strip())
            if search_project.strip():
                base_qs = base_qs.filter(project__project_id__icontains=search_project.strip())

            if pending_days_min > 0:
                from django.utils.timezone import now
                from datetime import timedelta
                cutoff = now() - timedelta(days=pending_days_min)
                base_qs = base_qs.filter(created_at__lte=cutoff)

            sort_map = {
                "Oldest first"  : "created_at",
                "Newest first"  : "-created_at",
                "Project ID"    : "project__project_id",
                "Team ID"       : "team__team_id",
            }
            pending_evals = base_qs.order_by(sort_map[sort_order])

            # ── Submission-date filter (needs per-row lookup so do after ORM) ──
            # Build submission date map for the current queryset
            eval_list = list(pending_evals)

            if date_from or date_to:
                from django.utils.timezone import make_aware
                from datetime import datetime, timezone as dt_tz
                filtered = []
                for ev in eval_list:
                    sd = get_submission_date(ev)
                    if sd is None:
                        continue
                    sd_date = sd.date()
                    if date_from and sd_date < date_from:
                        continue
                    if date_to and sd_date > date_to:
                        continue
                    filtered.append(ev)
                eval_list = filtered

            total = len(eval_list)

            if total == 0:
                st.info("No evaluations match the current filters.")
            else:
                # ── Metrics row ──────────────────────────────────────────────
                not_reviewed   = sum(1 for e in eval_list if not e.reviewed)
                rev_pending    = sum(1 for e in eval_list if e.reviewed and not e.email_sent)
                sheets_count   = sum(1 for e in eval_list if e.sheets_submitted)
                portal_count   = sum(1 for e in eval_list if not e.sheets_submitted)
                email_pending  = sum(1 for e in eval_list if not e.email_sent)

                mc1, mc2, mc3, mc4, mc5, mc6 = st.columns(6)
                mc1.metric("Total shown",           total)
                mc2.metric("Email pending",         email_pending)
                mc3.metric("Not yet reviewed",      not_reviewed)
                mc4.metric("Reviewed / unsent",     rev_pending)
                mc5.metric("🌐 Via portal",         portal_count)
                mc6.metric("📊 Via Google Sheets",  sheets_count)

                st.divider()

                # ── Table header ─────────────────────────────────────────────
                h1, h2, h3, h4, h5, h6, h7 = st.columns([2.2, 1.6, 0.8, 1.2, 1.1, 1.1, 0.9])
                h1.markdown("**Team ID**")
                h2.markdown("**Project**")
                h3.markdown("**Course**")
                h4.markdown("**Review**")
                h5.markdown("**Source**")
                h6.markdown("**Submitted**")
                h7.markdown("**Action**")
                st.divider()

                for ev in eval_list:
                    sub_date = get_submission_date(ev)
                    sub_str  = sub_date.strftime("%d %b %Y") if sub_date else "—"

                    review_badge  = "✅ Done"   if ev.reviewed        else "⏳ Pending"
                    source_badge  = "📊 Sheets" if ev.sheets_submitted else "🌐 Portal"

                    c1, c2, c3, c4, c5, c6, c7 = st.columns([2.2, 1.6, 0.8, 1.2, 1.1, 1.1, 0.9])
                    c1.write(ev.team.team_id)
                    c2.write(ev.project.project_id)
                    c3.write(ev.course)
                    c4.write(review_badge)
                    c5.write(source_badge)
                    c6.write(sub_str)
                    if c7.button("Open", key=f"btn_{ev.id}", use_container_width=True):
                        st.session_state.selected_eval_id = ev.id
                        st.rerun()

                st.divider()
        
        with tab2:
            st.markdown("""
            ### How to Review Evaluations
            1. Select an evaluation from the list above
            2. Read and edit the evaluation text as needed
            3. Update the grade if necessary
            4. Preview submission files for reference
            5. Click "Mark as Reviewed" to finalize

            ### File Formats
            - **Jupyter Notebooks (.ipynb)**: Preview in app
            - **ZIP Files (.zip)**: Interactive explorer
            - **Images**: Display in app
            - **Other formats**: Download for review
            """)

        if tab_admin:
            with tab_admin:
                from datetime import date, timedelta
                from django.utils.timezone import make_aware
                from datetime import datetime

                st.subheader("Trainer Review Activity")

                # ── Date range picker ─────────────────────────────────────────
                col_a, col_b, col_c = st.columns([1, 1, 2])
                with col_a:
                    from_date = st.date_input("From", value=date.today(), key="admin_from")
                with col_b:
                    to_date = st.date_input("To", value=date.today(), key="admin_to")

                if from_date > to_date:
                    st.warning("'From' date must be before 'To' date.")
                else:
                    from_dt = make_aware(datetime.combine(from_date, datetime.min.time()))
                    to_dt   = make_aware(datetime.combine(to_date,   datetime.max.time()))

                    # Reviews done in range, grouped by reviewer
                    reviewed_evals = (
                        Evaluation.objects
                        .filter(reviewed=True, reviewed_by__isnull=False,
                                reviewed_at__gte=from_dt, reviewed_at__lte=to_dt)
                        .values("reviewed_by", "reviewed_at__date")
                        .order_by("reviewed_by", "reviewed_at__date")
                    )

                    # Build summary: {email: {date: count}}
                    from collections import defaultdict
                    summary = defaultdict(lambda: defaultdict(int))
                    all_dates = set()
                    for row in reviewed_evals:
                        email = row["reviewed_by"]
                        d     = row["reviewed_at__date"]
                        summary[email][d] += 1
                        all_dates.add(d)

                    # Totals row
                    all_trainers = TrainerProfile.objects.filter(is_active=True, is_admin=False).order_by("name")

                    if not all_dates and not summary:
                        st.info("No reviews found for the selected date range.")
                    else:
                        sorted_dates = sorted(all_dates)

                        # ── Overall metrics ───────────────────────────────────
                        total_reviews = sum(
                            cnt for trainer_data in summary.values()
                            for cnt in trainer_data.values()
                        )
                        active_trainers = len(summary)

                        m1, m2, m3 = st.columns(3)
                        m1.metric("Total Reviews", total_reviews)
                        m2.metric("Active Trainers", active_trainers)
                        m3.metric("Days in Range", (to_date - from_date).days + 1)

                        st.divider()

                        # ── Per-trainer breakdown table ───────────────────────
                        st.markdown("**Reviews per trainer per day**")

                        # Header
                        date_cols = [str(d) for d in sorted_dates]
                        header_cols = st.columns([2] + [1] * len(sorted_dates) + [1])
                        header_cols[0].markdown("**Trainer**")
                        for i, d in enumerate(sorted_dates):
                            header_cols[i + 1].markdown(f"**{d.strftime('%d %b')}**")
                        header_cols[-1].markdown("**Total**")

                        st.divider()

                        # Rows — one per trainer
                        for trainer in all_trainers:
                            trainer_data = summary.get(trainer.email, {})
                            row_total = sum(trainer_data.values())
                            row_cols = st.columns([2] + [1] * len(sorted_dates) + [1])
                            row_cols[0].write(f"{trainer.name}")
                            for i, d in enumerate(sorted_dates):
                                count = trainer_data.get(d, 0)
                                row_cols[i + 1].write(str(count) if count else "—")
                            row_cols[-1].markdown(f"**{row_total}**")

                        st.divider()

                        # ── Edit log summary ──────────────────────────────────
                        st.markdown("**Edits made to LLM evaluations**")
                        edit_logs = (
                            TrainerEditLog.objects
                            .filter(edited_at__gte=from_dt, edited_at__lte=to_dt)
                            .select_related("evaluation__team", "evaluation__project")
                            .order_by("-edited_at")
                        )
                        if not edit_logs:
                            st.info("No edits in this period.")
                        else:
                            for log in edit_logs:
                                with st.expander(
                                    f"{log.trainer_email} — {log.evaluation.team.team_id} / "
                                    f"{log.evaluation.project.project_id} — "
                                    f"{log.edited_at.strftime('%d %b %Y %H:%M')}"
                                ):
                                    st.markdown(f"**Reason:** {log.feedback}")
                                    c1, c2 = st.columns(2)
                                    with c1:
                                        st.markdown("**Original**")
                                        st.text(log.original_text[:800] + ("…" if len(log.original_text) > 800 else ""))
                                    with c2:
                                        st.markdown("**Edited**")
                                        st.text(log.edited_text[:800] + ("…" if len(log.edited_text) > 800 else ""))
    
    if st.session_state.selected_eval_id:
        st.subheader("Review Evaluation Details")
        
        selected_eval = Evaluation.objects.filter(id=st.session_state.selected_eval_id).select_related("team", "project").first()
        
        if not selected_eval:
            st.error("Evaluation not found.")
            if st.button("Back to List"):
                st.session_state.selected_eval_id = None
                st.rerun()
        else:
            # Source banner
            if selected_eval.sheets_submitted:
                st.info("📊 **Submitted via Google Sheets** — response was sent outside the portal mailer")
            else:
                st.success("🌐 **Submitted via Portal** — evaluation processed through the portal")

            col1, col2, col3, col4 = st.columns(4)
            with col1:
                st.metric("Team ID", selected_eval.team.team_id)
            with col2:
                st.metric("Project ID", selected_eval.project.project_id)
            with col3:
                st.metric("Course", selected_eval.course)
            with col4:
                st.metric("Created", selected_eval.created_at.strftime("%b %d, %Y"))
            
            st.divider()
            
            eval_text = selected_eval.evaluation_report
            if eval_text:
                eval_text = _read_eval_file(eval_text)
            
            current_grade = selected_eval.evaluation_grade or ""

            eid = selected_eval.id
            if eid not in st.session_state.eval_text_edited:
                st.session_state.eval_text_edited[eid] = eval_text
            if eid not in st.session_state.eval_grade_edited:
                st.session_state.eval_grade_edited[eid] = current_grade
            # Store the original (LLM) text once — used to detect edits
            if eid not in st.session_state.eval_text_original:
                st.session_state.eval_text_original[eid] = eval_text
            if eid not in st.session_state.eval_edit_feedback:
                st.session_state.eval_edit_feedback[eid] = ""

            st.subheader("Edit Evaluation")
            edited_text = st.text_area(
                "Evaluation Report",
                value=st.session_state.eval_text_edited[eid],
                height=300,
                key=f"eval_text_{eid}",
            )
            st.session_state.eval_text_edited[eid] = edited_text

            edited_grade = st.text_input(
                "Grade",
                value=st.session_state.eval_grade_edited[eid],
                placeholder="e.g., A, B+, 85, 8.5/10",
                key=f"eval_grade_{eid}",
            )
            st.session_state.eval_grade_edited[eid] = edited_grade

            # ── Detect changes ────────────────────────────────────────────────
            text_was_edited = (
                edited_text is not None
                and st.session_state.eval_text_original.get(eid) is not None
                and edited_text.strip() != (st.session_state.eval_text_original[eid] or "").strip()
            )
            original_grade = selected_eval.evaluation_grade or ""
            grade_was_edited = edited_grade.strip() != original_grade.strip()

            # ── Mandatory comment only when something was changed ─────────────
            something_changed = text_was_edited or grade_was_edited
            if something_changed:
                st.divider()
                st.warning("You have made changes to this evaluation. A comment is required before submitting.")
                edit_feedback = st.text_area(
                    "Reason for changes (required)",
                    value=st.session_state.eval_edit_feedback[eid],
                    placeholder="Briefly explain what you changed and why — one sentence is enough…",
                    height=100,
                    key=f"edit_feedback_{eid}",
                )
                st.session_state.eval_edit_feedback[eid] = edit_feedback
            
            st.divider()
            st.subheader("Submission Files")
            
            from Submissions.models import Submission
            latest_submission = Submission.objects.filter(team_project__team=selected_eval.team, team_project__project=selected_eval.project).select_related("team_project").order_by("-created_at").first()
            
            if not latest_submission:
                st.info("No submission files found.")
            else:
                submission_files = latest_submission.files.all()
                if not submission_files:
                    st.info("No files attached to submission.")
                else:
                    from django.core.files.storage import default_storage
                    for sub_file in submission_files:
                        file_name = sub_file.original_name or sub_file.file.name
                        try:
                            file_data = _open_submission_file(sub_file.file.name)
                            file_ext = Path(file_name).suffix.lower()
                            col1, col2 = st.columns([3, 1])
                            with col1:
                                st.write(f"📄 {file_name}")
                            with col2:
                                if file_ext == '.ipynb':
                                    if st.button("Preview", key=f"view_{sub_file.id}"):
                                        st.session_state[f"show_notebook_{sub_file.id}"] = True
                                elif file_ext == '.zip':
                                    if st.button("Explore", key=f"explore_{sub_file.id}"):
                                        st.session_state[f"show_zip_{sub_file.id}"] = True
                                else:
                                    st.download_button(label="Download", data=file_data, file_name=file_name, key=f"dl_{sub_file.id}")
                        except Exception as e:
                            st.error(f"Error processing file {file_name}: {str(e)}")
                    # Display notebook preview
                    for sub_file in submission_files:
                        if sub_file.original_name and sub_file.original_name.endswith('.ipynb'):
                            show_key = f"show_notebook_{sub_file.id}"
                            if st.session_state.get(show_key, False):
                                file_name = sub_file.original_name or sub_file.file.name
                                try:
                                    nb_data = _open_submission_file(sub_file.file.name)
                                    import tempfile
                                    with tempfile.NamedTemporaryFile(suffix='.ipynb', delete=False, mode='wb') as tmp:
                                        tmp.write(nb_data)
                                        tmp.flush()
                                        if NBCONVERT_AVAILABLE:
                                            try:
                                                exporter = HTMLExporter()
                                                html_content, _ = exporter.from_filename(tmp.name)
                                                st.components.v1.html(html_content, height=800, scrolling=True)
                                            except Exception as e:
                                                st.error(f"Error rendering notebook: {str(e)}")
                                        else:
                                            st.warning("Notebook preview unavailable.")
                                    os.unlink(tmp.name)
                                    if st.button("Close Preview", key=f"close_{sub_file.id}"):
                                        st.session_state[show_key] = False
                                        st.rerun()
                                except Exception as e:
                                    st.error(f"Error reading notebook: {str(e)}")
                    # Display zip explorer
                    for sub_file in submission_files:
                        if sub_file.original_name and sub_file.original_name.endswith('.zip'):
                            show_key = f"show_zip_{sub_file.id}"
                            if st.session_state.get(show_key, False):
                                file_name = sub_file.original_name or sub_file.file.name
                                try:
                                    zip_data = _open_submission_file(sub_file.file.name)
                                    import tempfile
                                    with tempfile.NamedTemporaryFile(suffix='.zip', delete=False, mode='wb') as tmp:
                                        tmp.write(zip_data)
                                        tmp.flush()
                                        st.divider()
                                        st.subheader(f"ZIP Explorer: {file_name}")
                                        render_zip_explorer(tmp.name)
                                    os.unlink(tmp.name)
                                    if st.button("Close Explorer", key=f"close_zip_{sub_file.id}"):
                                        st.session_state[show_key] = False
                                        st.rerun()
                                except Exception as e:
                                    st.error(f"Error reading zip file: {str(e)}")
            
            st.divider()
            col1, col2 = st.columns(2)
            with col1:
                if st.button("Back to List", use_container_width=True):
                    st.session_state.selected_eval_id = None
                    st.rerun()

            with col2:
                submit_blocked = something_changed and not st.session_state.eval_edit_feedback.get(eid, "").strip()
                if submit_blocked:
                    st.button("Mark as Reviewed", type="primary", use_container_width=True, disabled=True)
                    st.caption("Fill in the reason above to enable submission.")
                elif st.button("Mark as Reviewed", type="primary", use_container_width=True, key=f"submit_{eid}"):
                    try:
                        new_eval_text = st.session_state.eval_text_edited[eid]
                        new_grade     = st.session_state.eval_grade_edited[eid]

                        # Save edit audit log if text or grade was changed
                        if text_was_edited or grade_was_edited:
                            TrainerEditLog.objects.create(
                                evaluation=selected_eval,
                                trainer_email=st.session_state.trainer_email,
                                original_text=st.session_state.eval_text_original.get(eid, ""),
                                edited_text=new_eval_text,
                                feedback=st.session_state.eval_edit_feedback.get(eid, ""),
                            )

                        # Update evaluation file (S3/local agnostic)
                        if selected_eval.evaluation_report:
                            from django.core.files.storage import default_storage
                            from django.core.files.base import ContentFile
                            try:
                                default_storage.delete(selected_eval.evaluation_report)
                                default_storage.save(
                                    selected_eval.evaluation_report,
                                    ContentFile(new_eval_text.encode('utf-8'))
                                )
                            except Exception as e:
                                st.error(f"Could not save evaluation file: {str(e)}")

                        selected_eval.evaluation_grade = new_grade
                        selected_eval.reviewed = True
                        selected_eval.reviewed_at = timezone.now()
                        selected_eval.reviewed_by = st.session_state.trainer_email
                        # Set visibility_after to now so the email scheduler picks it up immediately
                        if not selected_eval.visibility_after:
                            selected_eval.visibility_after = timezone.now()
                        selected_eval.save(update_fields=["evaluation_grade", "reviewed", "reviewed_at", "reviewed_by", "visibility_after"])

                        for _k in ["eval_text_edited", "eval_grade_edited", "eval_text_original", "eval_edit_feedback"]:
                            st.session_state[_k].pop(eid, None)

                        st.session_state.selected_eval_id = None
                        st.success("Evaluation marked as reviewed! Scheduled job will send email.")
                        st.rerun()

                    except Exception as e:
                        st.error(f"Error saving evaluation: {str(e)}")